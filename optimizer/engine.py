"""Izračun ABC/XYZ, sigurnosne zalihe, točke naručivanja i EOQ.

Sigurnosna zaliha koristi pojednostavljenu formulu bez varijabilnosti roka isporuke:
    SZ = z * sigma_D * sqrt(LT)
gdje je sigma_D standardna devijacija mjesečne potražnje, a LT rok isporuke
izražen u mjesecima (dani / DANA_U_MJESECU).
"""

from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass, field, asdict

import numpy as np
import pandas as pd

DANA_U_MJESECU = 30.0

Z_PO_RAZINI = {90: 1.28, 95: 1.65, 98: 2.05, 99: 2.33}

ALIASI = {
    "sifra": ["sifra", "šifra", "code", "sku", "artikl", "item"],
    "naziv": ["naziv", "name", "opis", "description"],
    "dobavljac": ["dobavljac", "dobavljač", "supplier", "vendor"],
    "nabavna": ["nabavna", "nabavna cijena", "nabavna_cijena", "cost", "unit cost"],
    "prodajna": ["prodajna", "prodajna cijena", "prodajna_cijena", "price", "unit price"],
    "zaliha": ["prosjecna zaliha", "prosječna zaliha", "prosjecna_zaliha", "zaliha", "avg stock", "stock"],
    "trenutna": ["trenutna zaliha", "trenutna_zaliha", "stanje", "na stanju", "on hand", "current stock"],
    "na_putu": ["na putu", "naruceno", "naručeno", "otvorene narudzbe", "otvorene narudžbe", "on order", "in transit"],
    "lead_time": ["lead time", "lead_time", "leadtime", "rok isporuke", "rok", "lt"],
    "moq": ["moq", "min narudzba", "min. narudžba", "minimalna narudzba"],
    "pakiranje": ["pakiranje", "pack", "pack size", "paket"],
}

OBAVEZNI = ["sifra", "naziv", "dobavljac", "nabavna", "prodajna", "zaliha"]
OZNAKE_UKUPNO = {"total", "ukupno", "sveukupno", "grand total", "sum", "zbroj"}


class GreskaUnosa(Exception):
    """Podaci kupca nisu upotrebljivi; poruka je namijenjena kupcu."""


@dataclass
class Parametri:
    razina_a: int = 98
    razina_b: int = 95
    razina_c: int = 90
    trosak_narudzbe: float = 60.0   # S, EUR po narudžbi
    drzanje_postotak: float = 20.0  # h, % nabavne cijene godišnje (H = h * nabavna)
    lead_time_default: float = 21.0  # dani, kad izvoz nema lead time
    cf_faktori: tuple = (25, 40, 60, 75, 85, 95, 100)  # % realne vrijednosti zalihe po pojasu obrtaja (procjena)
    pokrivenost_tjedana: float = 4.0  # količina po narudžbi kad nema nabavne cijene (EOQ nije moguć)
    abc_a: float = 80.0
    abc_b: float = 95.0
    xyz_x: float = 0.5
    xyz_y: float = 1.0

    def z_za(self, klasa: str) -> float:
        razina = {"A": self.razina_a, "B": self.razina_b, "C": self.razina_c}[klasa]
        return z_vrijednost(razina)


CF_POJASI = ["Bez prodaje", "do 1", "1 do 2", "2 do 3", "3 do 4", "4 do 6", "preko 6"]


def z_vrijednost(razina: float) -> float:
    """Z za razinu usluge u postocima; tablične vrijednosti, inače inverz normalne razdiobe."""
    r = int(round(razina))
    if r in Z_PO_RAZINI:
        return Z_PO_RAZINI[r]
    from statistics import NormalDist
    if not 50 < razina < 100:
        raise GreskaUnosa(f"Razina usluge {razina}% nije dopuštena (mora biti između 50 i 100).")
    return round(NormalDist().inv_cdf(razina / 100.0), 2)


@dataclass
class Rezultat:
    artikli: pd.DataFrame
    uskladjenost: pd.DataFrame
    sazetak: dict
    upozorenja: list[str] = field(default_factory=list)
    pretpostavke: list[str] = field(default_factory=list)


# ---------------------------------------------------------------- učitavanje

def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s).strip().lower())


def ucitaj(sadrzaj: bytes, ime_datoteke: str) -> pd.DataFrame:
    ime = ime_datoteke.lower()
    try:
        if ime.endswith((".xlsx", ".xlsm", ".xls")):
            return pd.read_excel(io.BytesIO(sadrzaj))
        if ime.endswith((".csv", ".txt")):
            for kodiranje in ("utf-8-sig", "cp1250"):  # UTF-8 (s BOM-om ili bez), pa Windows-1250 iz starijih ERP-ova
                try:
                    tekst = sadrzaj.decode(kodiranje)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                raise GreskaUnosa("Kodiranje CSV datoteke nije prepoznato. Spremite je kao UTF-8 ili učitajte Excel.")
            decimal = "," if ";" in tekst[:2000] else "."
            return pd.read_csv(io.StringIO(tekst), sep=None, engine="python", decimal=decimal)
    except GreskaUnosa:
        raise
    except Exception as e:  # noqa: BLE001
        raise GreskaUnosa(f"Datoteku nije moguće pročitati: {e}") from e
    raise GreskaUnosa("Podržani su formati .xlsx i .csv.")


def _num(s: pd.Series) -> pd.Series:
    """Pretvara stupac u brojeve; tekstualni stupci (npr. zbog opisnog retka) smiju imati decimalni zarez."""
    if s.dtype == object:
        s = s.astype(str).str.strip().str.replace(",", ".", regex=False)
    return pd.to_numeric(s, errors="coerce")


def _dvanaest_uzastopnih(svi: list, kandidati: list) -> list | None:
    """Prvi niz od točno 12 susjednih kandidat-stupaca (redoslijed kao u datoteci)."""
    skup = set(kandidati)
    niz: list = []
    nizovi: list[list] = []
    for c in svi:
        if c in skup:
            niz.append(c)
        else:
            if niz:
                nizovi.append(niz)
            niz = []
    if niz:
        nizovi.append(niz)
    for n in nizovi:
        if len(n) == 12:
            return n
    return None


def priprema(df: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """Mapira stupce, čisti podatke i vraća (standardizirani okvir, upozorenja)."""
    upoz: list[str] = []
    df = df.copy()
    df.columns = [str(c).strip().lstrip("\ufeff") for c in df.columns]
    mapa: dict[str, str] = {}
    for kanon, aliasi in ALIASI.items():
        for stupac in df.columns:
            if _norm(stupac) in aliasi and stupac not in mapa.values():
                mapa[kanon] = stupac
                break

    if "sifra" not in mapa:
        raise GreskaUnosa("U izvozu nedostaje stupac sa šifrom artikla (npr. 'Šifra').")

    # redci čija šifra počinje s '#' su opisi stupaca iz špranci i ne ulaze u analizu
    df = df[~df[mapa["sifra"]].astype(str).str.strip().str.startswith("#")]
    if len(df) == 0:
        raise GreskaUnosa("Datoteka nema nijednog retka s podacima. Popunite šprancu (jedan redak po artiklu) pa je učitajte ponovno.")

    kandidati = [c for c in df.columns if c not in mapa.values() and _norm(c) not in OZNAKE_UKUPNO
                 and _num(df[c]).notna().sum() > 0]
    mjeseci = _dvanaest_uzastopnih(list(df.columns), kandidati)
    if mjeseci is None:
        raise GreskaUnosa(
            f"Treba mi 12 uzastopnih stupaca s mjesečnom prodajom (količine), a pronašao sam {len(kandidati)} "
            f"numeričkih stupaca ({', '.join(map(str, kandidati[:15]))}). "
            "Mjeseci moraju stajati jedan do drugoga; uklonite ostale brojčane stupce koje ne koristim.")
    ignorirani = [c for c in kandidati if c not in mjeseci]
    if ignorirani:
        upoz.append("Ovi brojčani stupci nisu prepoznati kao mjeseci prodaje i ignorirani su: "
                    + ", ".join(map(str, ignorirani)) + ".")

    out = pd.DataFrame({k: df[v] for k, v in mapa.items()})
    for i, c in enumerate(mjeseci, 1):
        out[f"m{i}"] = _num(df[c])
    if "naziv" not in out:
        out["naziv"] = out["sifra"]
    if "dobavljac" not in out:
        out["dobavljac"] = "(svi artikli)"
    for c in ("nabavna", "prodajna", "zaliha"):
        if c not in out:
            out[c] = np.nan

    # retci "Total"
    maska_total = out["sifra"].astype(str).map(_norm).isin(OZNAKE_UKUPNO) | \
        out["naziv"].astype(str).map(_norm).isin(OZNAKE_UKUPNO)
    if maska_total.any():
        upoz.append(f"Uklonjeno redaka s ukupnim zbrojem: {int(maska_total.sum())}.")
        out = out[~maska_total]
    prazni = out["sifra"].isna() | (out["sifra"].astype(str).str.strip() == "")
    if prazni.any():
        upoz.append(f"Uklonjeno redaka bez šifre: {int(prazni.sum())}.")
        out = out[~prazni]
    if out.empty:
        raise GreskaUnosa("Nakon čišćenja nema nijednog artikla.")

    for c in ["nabavna", "prodajna", "zaliha"] + [f"m{i}" for i in range(1, 13)]:
        out[c] = _num(out[c])
    for c in ("lead_time", "moq", "pakiranje", "trenutna", "na_putu"):
        if c in out:
            out[c] = _num(out[c])

    mcols = [f"m{i}" for i in range(1, 13)]
    if out["m12"].fillna(0).sum() == 0:
        upoz.append("Zadnji mjesec nema nikakve prodaje. Provjerite je li izvoz završen s tekućim mjesecom.")
    neg = (out[mcols] < 0).sum().sum()
    if neg:
        upoz.append(f"Pronađeno je {int(neg)} negativnih mjesečnih prodaja (povrati ili storna). Tretirane su kao 0.")
    out[mcols] = out[mcols].clip(lower=0).fillna(0)

    for c in ("nabavna", "prodajna", "zaliha"):
        if c not in mapa:
            continue
        losi = out[c].isna() | (out[c] < 0)
        if losi.any():
            upoz.append(f"{int(losi.sum())} artikala nema ispravan podatak '{c}' i zamijenjeno je s 0.")
            out[c] = out[c].where(~losi, 0)
    out["dobavljac"] = out["dobavljac"].fillna("(bez dobavljača)").astype(str)
    out["sifra"] = out["sifra"].astype(str)
    if out["sifra"].duplicated().any():
        upoz.append(f"Šifre se ponavljaju u {int(out['sifra'].duplicated().sum())} redaka; svaki redak računa se zasebno.")
    return out.reset_index(drop=True), upoz


# ------------------------------------------------------------------ izračun

def _round_up(x: float, korak: float) -> float:
    return math.ceil(x / korak - 1e-9) * korak if korak and korak > 0 else x


def _abc_klase(vrijednost: pd.Series, a_prag: float, b_prag: float) -> pd.Series:
    """ABC klase po kumulativnom udjelu (A do a_prag %, B do b_prag %, ostalo C); bez vrijednosti je C."""
    v = vrijednost.fillna(0).clip(lower=0)
    ukupno = v.sum()
    if ukupno <= 0:
        return pd.Series("C", index=v.index)
    red = v.sort_values(ascending=False, kind="mergesort")
    kum_prije = (red.cumsum() - red) / ukupno * 100
    k = pd.Series(np.where(kum_prije < a_prag, "A", np.where(kum_prije < b_prag, "B", "C")), index=red.index).reindex(v.index)
    k[v <= 0] = "C"
    return k


def izracunaj(df: pd.DataFrame, p: Parametri, upozorenja: list[str] | None = None) -> Rezultat:
    upoz = list(upozorenja or [])
    pretp: list[str] = []
    mcols = [f"m{i}" for i in range(1, 13)]
    a = df.copy()

    a["potraznja_god"] = a[mcols].sum(axis=1)
    a["potraznja_mj"] = a["potraznja_god"] / 12.0
    a["sigma_mj"] = a[mcols].std(axis=1, ddof=1)
    a["cv"] = np.where(a["potraznja_mj"] > 0, a["sigma_mj"] / a["potraznja_mj"].replace(0, np.nan), np.nan)
    # osnova za ABC: prodajna cijena, inače nabavna, inače količina
    nab = a["nabavna"].where(a["nabavna"] > 0)
    prod = a["prodajna"].where(a["prodajna"] > 0)
    ima_cijene = bool(nab.notna().any())
    ima_prodajnu = bool(prod.notna().any())
    if ima_prodajnu or ima_cijene:
        cijena = prod.fillna(nab).fillna(prod.median() if ima_prodajnu else nab.median())
        abc_osnova = "godišnji promet u €" if ima_prodajnu else "godišnja vrijednost po nabavnoj cijeni"
    else:
        cijena = pd.Series(1.0, index=a.index)
        abc_osnova = "godišnja količina (nema cijena u izvozu)"
        upoz.append("Izvoz nema cijene, pa se ABC računa po količini, a ne po vrijednosti. Rezultat je grublji: "
                    "jeftin artikl s velikom količinom može ispasti A.")
    a["promet"] = a["potraznja_god"] * cijena

    # ABC po godišnjem prometu, kumulativno (a uz njega i po marži i po količini)
    if a["promet"].sum() <= 0:
        raise GreskaUnosa("Ukupna prodaja je nula; provjerite mjesečne količine.")
    a["abc"] = _abc_klase(a["promet"], p.abc_a, p.abc_b)
    a["abc_kolicina"] = _abc_klase(a["potraznja_god"], p.abc_a, p.abc_b)

    # marža po komadu = prodajna − nabavna cijena (samo kad postoje obje)
    oba = (a["prodajna"] > 0) & (a["nabavna"] > 0)
    ima_marzu = bool(oba.any())
    a["marza_jed"] = (a["prodajna"] - a["nabavna"]).where(oba)
    a["marza_pct"] = (a["marza_jed"] / a["prodajna"] * 100).where(oba)
    a["marza_god"] = a["potraznja_god"] * a["marza_jed"]
    a["abc_marza"] = _abc_klase(a["marza_god"], p.abc_a, p.abc_b).where(oba) if ima_marzu else None
    if ima_marzu and (a["marza_god"] < 0).any():
        upoz.append(f"{int((a['marza_god'] < 0).sum())} artikala se prodaje ispod nabavne cijene; u ABC analizi po marži svrstani su u C.")

    a["xyz"] = np.where(a["potraznja_god"] <= 0, "Z",
                        np.where(a["cv"] <= p.xyz_x, "X", np.where(a["cv"] <= p.xyz_y, "Y", "Z")))
    a["klasa"] = a["abc"] + a["xyz"]

    # lead time
    if "lead_time" in a:
        nema_lt = a["lead_time"].isna() | (a["lead_time"] <= 0)
        a["lt_izvor"] = np.where(nema_lt, "default", "izvoz")
        a["lt_dani"] = a["lead_time"].where(~nema_lt, p.lead_time_default)
        if nema_lt.any():
            upoz.append(f"{int(nema_lt.sum())} artikala nema lead time; korišten je zadani rok od {p.lead_time_default:g} d.")
        mjereni = a.loc[~nema_lt, "lt_dani"]
        if len(mjereni) > 1 and mjereni.nunique() == 1:
            upoz.append(f"Svi artikli imaju isti lead time ({mjereni.iloc[0]:g} dana). Je li to stvarni ugovoreni rok "
                        "ili privremena vrijednost? Ako varira, sigurnosne su zalihe podcijenjene.")
    else:
        a["lt_dani"] = p.lead_time_default
        a["lt_izvor"] = "default"
        upoz.append(f"Izvoz nema lead time; za sve artikle korišten je zadani rok od {p.lead_time_default:g} d.")

    a["z"] = a["abc"].map(p.z_za)
    a["sigurnosna"] = a["z"] * a["sigma_mj"].fillna(0) * np.sqrt(a["lt_dani"] / DANA_U_MJESECU)
    a["potraznja_dan"] = a["potraznja_mj"] / DANA_U_MJESECU
    a["tocka_naruc"] = a["potraznja_dan"] * a["lt_dani"] + a["sigurnosna"]

    if a["zaliha"].isna().all() and "trenutna" in a and a["trenutna"].notna().any():
        a["zaliha"] = a["trenutna"]
        upoz.append("Izvoz nema prosječnu zalihu; za obrtaj i status korištena je trenutna zaliha, pa su ti pokazatelji približni.")
    ima_zalihu = bool(a["zaliha"].notna().any())

    # EOQ; bez nabavne cijene nije moguć pa se količina po narudžbi računa kao pokrivenost potražnje
    H = (p.drzanje_postotak / 100.0) * nab
    eoq = np.sqrt(2.0 * a["potraznja_god"] * p.trosak_narudzbe / H)
    pokrivenost = a["potraznja_god"] / 52.0 * p.pokrivenost_tjedana
    a["eoq_izvor"] = np.where(eoq.notna(), "EOQ", "pokrivenost")
    a["eoq_osnovni"] = eoq.fillna(pokrivenost).fillna(0)
    if not ima_cijene:
        upoz.append(f"Izvoz nema nabavnu cijenu, pa EOQ nije moguć. Količina po narudžbi je potražnja za "
                    f"{p.pokrivenost_tjedana:g} tjedna. Dodajte nabavnu cijenu za pravi EOQ i vrijednosti u eurima.")
    elif (a["eoq_izvor"] == "pokrivenost").any():
        upoz.append(f"{int((a['eoq_izvor'] == 'pokrivenost').sum())} artikala nema nabavnu cijenu; za njih je količina po "
                    f"narudžbi potražnja za {p.pokrivenost_tjedana:g} tjedna.")
    q = a["eoq_osnovni"].copy()
    if "moq" in a:
        q = np.maximum(q, a["moq"].fillna(0))
    a["eoq"] = q
    if "pakiranje" in a:
        a["eoq"] = [_round_up(x, k) if k and not pd.isna(k) else x for x, k in zip(a["eoq"], a["pakiranje"])]
    a["eoq"] = np.where(a["potraznja_god"] > 0, np.ceil(a["eoq"]), 0)
    a["narudzbi_god"] = np.where(a["eoq"] > 0, a["potraznja_god"] / a["eoq"].replace(0, np.nan), 0)

    # količina za narudžbu: (s, Q) pravilo nad zalihom na stanju + na putu
    if "trenutna" in a:
        na_putu = a["na_putu"].fillna(0) if "na_putu" in a else 0.0
        a["pozicija"] = a["trenutna"].fillna(0) + na_putu
        manjak = a["tocka_naruc"] - a["pozicija"]
        treba = (a["potraznja_god"] > 0) & (manjak >= 0) & (a["eoq"] > 0)
        n = np.where(treba, np.maximum(1, np.floor(manjak / a["eoq"].replace(0, np.nan)).fillna(0) + 1), 0)
        a["narudzba"] = n * a["eoq"]
        a["vrijednost_narudzbe"] = a["narudzba"] * nab
        if "na_putu" not in a:
            upoz.append("Izvoz nema stupac 'Na putu'. Količine za narudžbu ne uzimaju u obzir već naručenu robu.")
    else:
        a["narudzba"] = np.nan
        a["vrijednost_narudzbe"] = np.nan
        upoz.append("Izvoz nema stupac 'Trenutna zaliha', pa ne mogu odrediti što treba naručiti ovaj tjedan. "
                    "Prikazani su samo sigurnosna zaliha, točka naručivanja i količina po narudžbi.")

    # obrtaj (nabavna vrijednost prodaje / prosječna zaliha po nabavnoj = D / prosječna zaliha)
    a["obrtaj"] = np.where(a["zaliha"] > 0, a["potraznja_god"] / a["zaliha"].replace(0, np.nan), np.nan)
    a["dani_zalihe"] = np.where(a["potraznja_dan"] > 0, a["zaliha"] / a["potraznja_dan"].replace(0, np.nan), np.nan)
    a["vrijednost_zalihe"] = a["zaliha"] * nab

    # GMROI = godišnja bruto marža / prosječna zaliha po nabavnoj cijeni
    a["gmroi"] = (a["marza_god"] / a["vrijednost_zalihe"]).where(a["vrijednost_zalihe"] > 0)

    # cash-flow potencijal: vrijednost zalihe umanjena faktorom realizacije prema pojasu obrtaja
    ima_cf = bool(ima_cijene and ima_zalihu)
    ima_vr = a["vrijednost_zalihe"] > 0
    o = a["obrtaj"]
    pojas = np.select(
        [a["potraznja_god"] <= 0, o <= 1, o <= 2, o <= 3, o <= 4, o <= 6],
        [0, 1, 2, 3, 4, 5], default=6)
    a["cf_pojas"] = pd.Series(pojas, index=a.index).where(ima_vr)
    faktori = np.array(p.cf_faktori, dtype=float)
    a["cf_faktor"] = a["cf_pojas"].map(lambda i: faktori[int(i)] if pd.notna(i) else np.nan)
    a["cf_realna"] = a["vrijednost_zalihe"] * a["cf_faktor"] / 100

    # status (procjena na temelju prosječne zalihe, ne trenutnog stanja)
    gornja = a["sigurnosna"] + a["eoq"]

    def status(r_zaliha, r_prodaja, r_sz, r_gornja):
        if pd.isna(r_zaliha):
            return "Bez podatka o zalihi"
        if r_prodaja <= 0:
            return "Nema prodaje" if r_zaliha > 0 else "Neaktivan"
        if r_zaliha < r_sz:
            return "Rizik manjka"
        if r_zaliha <= r_gornja:
            return "Uredno"
        if r_zaliha <= 2 * r_gornja:
            return "Višak"
        return "Ozbiljan višak"

    a["status"] = [status(z, d, s, g) for z, d, s, g in
                   zip(a["zaliha"], a["potraznja_god"], a["sigurnosna"], gornja)]

    uskl = _uskladjenost(a, ima_zalihu, ima_cijene)
    saz = _sazetak(a, uskl, p, ima_cijene, ima_zalihu, abc_osnova)
    saz["ima_marzu"] = ima_marzu
    saz["cashflow"] = _cashflow(a, p) if ima_cf else None
    gm_pod = a["gmroi"].notna()
    saz["gmroi"] = (float(a.loc[gm_pod, "marza_god"].sum() / a.loc[gm_pod, "vrijednost_zalihe"].sum())
                    if gm_pod.any() and a.loc[gm_pod, "vrijednost_zalihe"].sum() > 0 else None)
    saz["marza_ukupno"] = float(a["marza_god"].sum()) if ima_marzu else None

    pretp += [
        f"Sigurnosna zaliha = z × σ potražnje × √LT (bez varijabilnosti roka isporuke); z po klasi: "
        f"A {p.razina_a}% ({p.z_za('A')}), B {p.razina_b}% ({p.z_za('B')}), C {p.razina_c}% ({p.z_za('C')}).",
        "Ova sigurnosna zaliha pretpostavlja da rok isporuke ne varira. To je donja granica.",
        (f"EOQ = √(2 × godišnja potražnja × S / H), uz S = {p.trosak_narudzbe:g} € po narudžbi i "
         f"H = {p.drzanje_postotak:g}% nabavne cijene godišnje.") if ima_cijene else
        f"Količina po narudžbi = potražnja za {p.pokrivenost_tjedana:g} tjedna (EOQ zahtijeva nabavnu cijenu).",
        f"ABC po osnovi: {abc_osnova} (A do {p.abc_a:g}%, B do {p.abc_b:g}% kumulativno); "
        f"XYZ po koeficijentu varijacije (X ≤ {p.xyz_x:g}, Y ≤ {p.xyz_y:g}).",
        f"Mjesec se računa kao {DANA_U_MJESECU:g} dana; σ potražnje je uzorkovana standardna devijacija 12 mjeseci.",
        "Narudžba: ako je zaliha na stanju plus na putu na ili ispod točke naručivanja, naručuje se višekratnik EOQ-a "
        "(zaokruženo na MOQ i pakiranje) dovoljan da pozicija prijeđe točku naručivanja.",
        *([f"Cash-flow potencijal: vrijednost zalihe po nabavnoj cijeni umanjena faktorom realizacije prema koeficijentu obrtaja "
           f"({', '.join(f'{l}: {f:g}%' for l, f in zip(CF_POJASI, p.cf_faktori))}). Faktori su procjena i mijenjaju se u pretpostavkama."]
          if ima_cf else []),
        *(["GMROI = godišnja bruto marža / prosječna zaliha po nabavnoj cijeni: koliko eura bruto marže godišnje donosi euro uložen u zalihe."]
          if ima_marzu and ima_zalihu else []),
        "Status zalihe uspoređuje prosječnu zalihu sa sigurnosnom zalihom i EOQ-om; to je procjena, ne trenutno stanje.",
    ]
    return Rezultat(a, uskl, saz, upoz, pretp)


def _uskladjenost(a: pd.DataFrame, ima_zalihu: bool = True, ima_cijene: bool = True) -> pd.DataFrame:
    stupci = ["polje", "artikala", "udio_prodaja", "udio_zaliha", "odstupanje", "nalaz", "smjer", "rizik_stockout_A"]
    if not ima_zalihu:
        return pd.DataFrame(columns=stupci)
    tezina = a["zaliha"].fillna(0) * (a["nabavna"].where(a["nabavna"] > 0).fillna(a["nabavna"].median()) if ima_cijene else 1.0)
    ukupno_prodaja = a["promet"].sum()
    ukupno_zaliha = tezina.sum()
    redovi = []
    for abc in "ABC":
        for xyz in "XYZ":
            m = (a["abc"] == abc) & (a["xyz"] == xyz)
            up = a.loc[m, "promet"].sum() / ukupno_prodaja * 100 if ukupno_prodaja else 0.0
            uz = tezina[m].sum() / ukupno_zaliha * 100 if ukupno_zaliha else 0.0
            odst = uz - up
            nalaz = "uredno"
            if abs(odst) >= 5:
                nalaz = "nalaz"
            elif abs(odst) >= 2:
                nalaz = "spomenuti"
            rizik = False
            if abc == "A" and up > 0 and uz < 0.8 * up:
                rizik = True
                nalaz = "nalaz"
            smjer = "ok"
            if rizik or odst <= -2:
                smjer = "stockout"
            elif odst >= 2:
                smjer = "kapital"
            redovi.append({"polje": abc + xyz, "artikala": int(m.sum()),
                           "udio_prodaja": up, "udio_zaliha": uz, "odstupanje": odst,
                           "nalaz": nalaz, "smjer": smjer, "rizik_stockout_A": rizik})
    return pd.DataFrame(redovi)


def _cashflow(a: pd.DataFrame, p: Parametri) -> dict:
    ima = a["cf_pojas"].notna()
    ukupno = float(a.loc[ima, "vrijednost_zalihe"].sum())
    redovi = []
    for i, naziv in enumerate(CF_POJASI):
        m = ima & (a["cf_pojas"] == i)
        v = float(a.loc[m, "vrijednost_zalihe"].sum())
        f = float(p.cf_faktori[i])
        redovi.append({"pojas": naziv, "artikala": int(m.sum()), "vrijednost": v,
                       "udio": v / ukupno * 100 if ukupno else 0.0, "faktor": f,
                       "realna": v * f / 100, "umanjenje": v - v * f / 100})
    realna = sum(r["realna"] for r in redovi)
    return {"redovi": redovi, "vrijednost": ukupno, "realna": realna, "umanjenje": ukupno - realna}


def hr(x: float) -> str:
    return f"{x:.1f}".replace(".", ",")


def _sazetak(a: pd.DataFrame, uskl: pd.DataFrame, p: Parametri, ima_cijene: bool = True,
             ima_zalihu: bool = True, abc_osnova: str = "") -> dict:
    nab = a["nabavna"].where(a["nabavna"] > 0)
    vz = float(a["vrijednost_zalihe"].sum()) if ima_cijene and ima_zalihu else None
    if ima_zalihu:
        if ima_cijene:
            obrtaj = float((a["potraznja_god"] * nab).sum()) / vz if vz else None
        else:
            z = float(a["zaliha"].sum())
            obrtaj = float(a["potraznja_god"].sum()) / z if z > 0 else None
    else:
        obrtaj = None
    sz_vr = float((a["sigurnosna"] * nab).sum()) if ima_cijene else None
    kapital = uskl[uskl["smjer"] == "kapital"].sort_values("odstupanje", ascending=False)
    stock = uskl[uskl["smjer"] == "stockout"].sort_values("odstupanje")
    r = []
    if not stock.empty and stock.iloc[0]["nalaz"] != "uredno":
        x = stock.iloc[0]
        uvod = "Rizik manjka je" if x["nalaz"] == "nalaz" else "Blago manje zaliha nego prodaje je"
        r.append(f"{uvod} u polju {x['polje']}: čini {hr(x['udio_prodaja'])}% prodaje, "
                 f"a drži {hr(x['udio_zaliha'])}% vrijednosti zaliha.")
    if not kapital.empty and kapital.iloc[0]["nalaz"] != "uredno":
        k = kapital.iloc[0]
        uvod = "Novac stoji" if k["nalaz"] == "nalaz" else "Blago više zaliha nego prodaje je"
        r.append(f"{uvod} u polju {k['polje']}: nosi {hr(k['udio_prodaja'])}% prodaje, "
                 f"a veže {hr(k['udio_zaliha'])}% vrijednosti zaliha.")
    if not ima_zalihu:
        r.append("Za usporedbu strukture zaliha i prodaje treba stupac Prosječna zaliha (ili Trenutna zaliha).")
    elif not r:
        r.append("Struktura zaliha uglavnom prati strukturu prodaje; nijedno polje ne odstupa više od 2 postotna boda.")
    return {
        "broj_artikala": int(len(a)),
        "godisnji_promet": float(a["promet"].sum()),
        "vrijednost_zaliha": vz,
        "obrtaj": obrtaj,
        "dani_zalihe": 365.0 / obrtaj if obrtaj else None,
        "vrijednost_sigurnosnih": sz_vr,
        "klase_abc": {k: int(v) for k, v in a["abc"].value_counts().items()},
        "klase_xyz": {k: int(v) for k, v in a["xyz"].value_counts().items()},
        "statusi": {k: int(v) for k, v in a["status"].value_counts().items()},
        "narudzba_stavki": int((a["narudzba"] > 0).sum()) if a["narudzba"].notna().any() else None,
        "narudzba_vrijednost": float(a["vrijednost_narudzbe"].sum()) if a["narudzba"].notna().any() and ima_cijene else None,
        "vrijednost_visak": float(a.loc[a["status"].isin(["Višak", "Ozbiljan višak", "Nema prodaje"]),
                                        "vrijednost_zalihe"].sum()) if ima_cijene and ima_zalihu else None,
        "ima_cijene": ima_cijene, "ima_zalihu": ima_zalihu, "abc_osnova": abc_osnova,
        "recenice": r,
    }


def po_dobavljacu(a: pd.DataFrame) -> pd.DataFrame:
    nab = a["nabavna"].where(a["nabavna"] > 0)
    zbroj = lambda s: s.sum(min_count=1)  # noqa: E731  (NaN, a ne 0, kad cijena ne postoji)
    g = a.assign(vrijednost_eoq=a["eoq"] * nab, vr_sig=a["sigurnosna"] * nab, kol=a["eoq"]).groupby("dobavljac").agg(
        artikala=("sifra", "count"),
        vrijednost_eoq=("vrijednost_eoq", zbroj),
        vrijednost_sigurnosnih=("vr_sig", zbroj),
        vrijednost_zalihe=("vrijednost_zalihe", zbroj),
        stavki_narudzbe=("narudzba", lambda s: int((s > 0).sum())),
        vrijednost_narudzbe=("vrijednost_narudzbe", zbroj),
        kolicina_narudzbe=("narudzba", zbroj),
    ).reset_index().sort_values(["stavki_narudzbe", "artikala"], ascending=False)
    return g


def narudzba_csv(a: pd.DataFrame) -> pd.DataFrame:
    """Stavke za naručivanje, spremne za uvoz u ERP: jedan redak po artiklu.

    Stupci s cijenom i vrijednošću izostavljaju se kad izvoz nema nabavnu cijenu.
    """
    ima_cijene = bool((a["nabavna"] > 0).any())
    osnovni = ["dobavljac", "sifra", "naziv", "kolicina"]
    if a["narudzba"].isna().all():
        return pd.DataFrame(columns=osnovni + (["nabavna_cijena", "vrijednost"] if ima_cijene else []))
    n = a[a["narudzba"] > 0].copy()
    n["vr"] = n["vrijednost_narudzbe"].fillna(0)
    n = n.sort_values(["dobavljac", "vr", "narudzba"], ascending=[True, False, False])
    izlaz = pd.DataFrame({"dobavljac": n["dobavljac"], "sifra": n["sifra"], "naziv": n["naziv"],
                          "kolicina": n["narudzba"].astype(int)})
    if ima_cijene:
        izlaz["nabavna_cijena"] = n["nabavna"].round(2)
        izlaz["vrijednost"] = n["vrijednost_narudzbe"].round(2)
    return izlaz


def izvezi_excel(rez: Rezultat, p: Parametri) -> bytes:
    a = rez.artikli
    stupci = {
        "sifra": "Šifra", "naziv": "Naziv", "dobavljac": "Dobavljač", "klasa": "ABC/XYZ",
        "potraznja_god": "Potražnja god.", "sigma_mj": "σ potražnje (mj.)", "cv": "CV",
        "lt_dani": "Lead time (dani)", "lt_izvor": "Izvor lead timea", "z": "z",
        "sigurnosna": "Sigurnosna zaliha", "tocka_naruc": "Točka naručivanja", "eoq": "Količina po narudžbi", "eoq_izvor": "Način izračuna količine",
        "narudzbi_god": "Narudžbi god.", "trenutna": "Trenutna zaliha", "na_putu": "Na putu", "narudzba": "Za naručiti", "zaliha": "Prosječna zaliha", "obrtaj": "Obrtaj",
        "dani_zalihe": "Dani zalihe", "vrijednost_zalihe": "Vrijednost zalihe (€)", "marza_pct": "Marža %", "gmroi": "GMROI",
        "abc_marza": "ABC po marži", "cf_realna": "Realna vrijednost (€)", "status": "Status",
    }
    stupci = {k: v for k, v in stupci.items() if k in a}
    matrica = a[list(stupci)].rename(columns=stupci).copy()
    for c in ("σ potražnje (mj.)", "CV", "Sigurnosna zaliha", "Točka naručivanja", "Narudžbi god.",
              "Obrtaj", "Dani zalihe", "Vrijednost zalihe (€)"):
        matrica[c] = matrica[c].round(2)
    matrica["Sigurnosna zaliha"] = np.ceil(a["sigurnosna"]).astype(int)
    matrica["Točka naručivanja"] = np.ceil(a["tocka_naruc"]).astype(int)
    u = rez.uskladjenost.rename(columns={
        "polje": "Polje", "artikala": "Artikala", "udio_prodaja": "Udio u prodaji (%)",
        "udio_zaliha": "Udio u zalihama (%)", "odstupanje": "Odstupanje (p.b.)",
        "nalaz": "Nalaz", "smjer": "Smjer", "rizik_stockout_A": "Rizik stockouta A"}).round(2)
    d = po_dobavljacu(a).rename(columns={
        "dobavljac": "Dobavljač", "artikala": "Artikala", "stavki_narudzbe": "Stavki za naručiti",
        "vrijednost_narudzbe": "Vrijednost narudžbe (€)", "vrijednost_eoq": "Vrijednost EOQ (€)",
        "vrijednost_sigurnosnih": "Vrijednost sig. zaliha (€)", "vrijednost_zalihe": "Vrijednost zalihe (€)"}).round(2)
    par = pd.DataFrame({"Pretpostavke": rez.pretpostavke + [""] + ["Upozorenja:"] + rez.upozorenja})

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        narudzba_csv(a).to_excel(w, sheet_name="Narudžba", index=False)
        cf = rez.sazetak.get("cashflow")
        if cf:
            tab = pd.DataFrame(cf["redovi"]).rename(columns={
                "pojas": "Koeficijent obrtaja", "artikala": "Artikala", "vrijednost": "Vrijednost zalihe (€)",
                "udio": "Udio (%)", "faktor": "Faktor realizacije (%)", "realna": "Realna vrijednost (€)", "umanjenje": "Umanjenje (€)"}).round(2)
            tab.loc[len(tab)] = ["Ukupno", int(tab["Artikala"].sum()), round(cf["vrijednost"], 2), 100.0, None,
                                 round(cf["realna"], 2), round(cf["umanjenje"], 2)]
            tab.to_excel(w, sheet_name="Cash-flow potencijal", index=False)
        matrica.to_excel(w, sheet_name="Matrica zaliha", index=False)
        u.to_excel(w, sheet_name="Usklađenost", index=False)
        d.to_excel(w, sheet_name="Po dobavljaču", index=False)
        par.to_excel(w, sheet_name="Na čemu počiva", index=False)
        for ws in w.book.worksheets:
            ws.freeze_panes = "A2"
            for col in ws.columns:
                sirina = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, sirina + 2), 90)
    return buf.getvalue()


def demo_podaci(n: int = 40, seed: int = 7) -> pd.DataFrame:
    """Izmišljeni, ali realistični podaci za predložak i demonstraciju."""
    rng = np.random.default_rng(seed)
    dobavljaci = ["Alfa d.o.o.", "Beta trgovina", "Gama komponente", "Delta uvoz"]
    redovi = []
    for i in range(n):
        baza = float(rng.choice([400, 150, 60, 20, 8], p=[.08, .17, .3, .3, .15]))
        nestalnost = float(rng.choice([0.15, 0.4, 0.9], p=[.35, .4, .25]))
        mj = np.clip(rng.normal(baza, baza * nestalnost, 12), 0, None).round()
        nab = round(float(rng.uniform(2, 60)), 2)
        pros = float(mj.mean())
        lt = int(rng.choice([7, 14, 21, 35]))
        r = {"Šifra": f"ART-{i + 1:03d}", "Naziv": f"Artikl {i + 1}", "Dobavljač": dobavljaci[i % 4],
             "Nabavna cijena": nab, "Prodajna cijena": round(nab * float(rng.uniform(1.25, 1.7)), 2),
             # većina artikala ima 0,5 do 3 mjeseca zalihe, nešto sporih ima i 6 do 30 mjeseci
             "Prosječna zaliha": int(pros * (rng.uniform(0.5, 3) if rng.random() < 0.8 else rng.uniform(6, 30))),
             "Trenutna zaliha": int(pros * rng.uniform(0, 3)), "Na putu": int(pros * rng.choice([0, 0, 1])),
             "Lead time": lt, "MOQ": int(rng.choice([0, 0, 20, 50])), "Pakiranje": int(rng.choice([1, 1, 6, 12]))}
        if i in (11, 29):  # dva artikla bez ikakve prodaje, ali sa zalihom
            mj = np.zeros(12)
            r["Prosječna zaliha"] = int(rng.integers(40, 200))
            r["Trenutna zaliha"] = r["Prosječna zaliha"]
        for k, v in enumerate(mj, 1):
            r[f"Mjesec {k}"] = int(v)
        redovi.append(r)
    return pd.DataFrame(redovi)


# Jedan izvor opisa stupaca: koristi se u špranci (komentar na zaglavlju i opisni redak) i u sučelju.
# (naziv, obavezno, što se dobiva kad se stupac popuni)
OPISI_STUPACA = [
    ("Šifra", True, "Jedinstvena oznaka artikla, jedan redak po artiklu. Bez nje nema analize."),
    ("Naziv", False, "Samo za čitljivost popisa artikala i narudžbe."),
    ("Dobavljač", False, "Narudžba razdvojena po dobavljaču i filtriranje cijele analize po dobavljaču."),
    ("Nabavna cijena", False, "Vrijednosti u eurima, pravi EOQ (optimalna količina narudžbe), obrtaj po vrijednosti i cash-flow potencijal zaliha."),
    ("Prodajna cijena", False, "Uz nabavnu cijenu: marža po artiklu, ABC analiza po marži i GMROI, tj. koliko eura bruto marže godišnje donosi euro uložen u zalihe."),
    ("Prosječna zaliha", False, "Matrica zaliha naspram prodaje, koeficijent obrtaja, dani zalihe, status artikla, cash-flow potencijal i GMROI."),
    ("Trenutna zaliha", False, "Količina koju treba naručiti ovaj tjedan."),
    ("Na putu", False, "Da se već naručena roba ne naruči dvaput."),
    ("Lead time", False, "Stvarni rok isporuke u danima umjesto zadanih 21: točnija sigurnosna zaliha i točka naručivanja."),
    ("MOQ", False, "Narudžba ne pada ispod minimalne količine dobavljača."),
    ("Pakiranje", False, "Narudžba se zaokružuje na cijelo pakiranje."),
    ("Mjesec 1 do 12", True, "Količine prodaje po mjesecima, zadnjih 12 mjeseci (Mjesec 1 je najstariji). Daju ABC/XYZ analizu, sigurnosnu zalihu, točku naručivanja i količinu po narudžbi."),
]
STUPCI_PREDLOSKA = [n for n, _, _ in OPISI_STUPACA[:-1]] + [f"Mjesec {i}" for i in range(1, 13)]
_OBAVEZNI_PREDLOSKA = {"Šifra"} | {f"Mjesec {i}" for i in range(1, 13)}


def _opis_stupca(naziv: str) -> tuple[bool, str]:
    for n, obav, tekst in OPISI_STUPACA:
        if n == naziv or (n == "Mjesec 1 do 12" and naziv.startswith("Mjesec ")):
            return obav, tekst
    return False, ""


def _opisni_redak() -> list[str]:
    """Drugi redak špranci: kratak opis što se dobiva uz svaki stupac. Šifra počinje s '#', pa ga učitavanje preskače."""
    redak = []
    for naziv in STUPCI_PREDLOSKA:
        obav, tekst = _opis_stupca(naziv)
        kratko = tekst.split(". ")[0].rstrip(".").replace(";", ",")
        redak.append(("OBAVEZNO. " if obav else "Neobavezno. ") + kratko)
    redak[0] = "# " + redak[0]
    return redak


_UPUTE_PREDLOSKA = [
    ("Što je obavezno", "Samo šifra artikla i 12 mjeseci prodaje (količine u komadima). Sve ostalo je dodatak koji otključava dodatne rezultate."),
    ("Drugi redak (opisi)", "Redak ispod zaglavlja objašnjava što dobivate uz svaki stupac. Počinje znakom #, pa ga aplikacija preskače: možete ga ostaviti ili obrisati."),
    ("Mjeseci", "Mjesec 1 je najstariji, Mjesec 12 najnoviji (zadnji zaključeni mjesec). Stupci moraju stajati jedan do drugoga."),
    ("Prazne ćelije", "Prazna ćelija prodaje računa se kao 0. Negativne prodaje (povrati) računaju se kao 0."),
    ("Retci", "Jedan redak po artiklu. Obrišite retke s ukupnim zbrojem (Total)."),
    ("Nazivi stupaca", "Možete mijenjati redoslijed stupaca i brisati neobavezne. Aplikacija prepoznaje i uobičajene nazive (npr. SKU, Supplier, Stanje)."),
    ("Količine i cijene", "Količine su u komadima (ili osnovnoj jedinici mjere), cijene u EUR po komadu."),
]


def predlozak_excel(s_primjerom: bool = False) -> bytes:
    """Prazna (ili popunjena izmišljenim podacima) špranca za unos izvoza iz ERP-a, s opisnim retkom."""
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    if s_primjerom:
        podaci = demo_podaci()[STUPCI_PREDLOSKA]
    else:
        podaci = pd.DataFrame(columns=STUPCI_PREDLOSKA)
    opis = pd.DataFrame([_opisni_redak()], columns=STUPCI_PREDLOSKA)
    podaci = pd.concat([opis, podaci], ignore_index=True)
    upute = pd.DataFrame(_UPUTE_PREDLOSKA, columns=["Tema", "Objašnjenje"])
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        podaci.to_excel(w, sheet_name="Podaci", index=False)
        upute.to_excel(w, sheet_name="Upute", index=False)
        ws = w.book["Podaci"]
        for i, naziv in enumerate(STUPCI_PREDLOSKA, 1):
            c = ws.cell(row=1, column=i)
            obav, tekst = _opis_stupca(naziv)
            c.fill = PatternFill("solid", fgColor="FF8A3D" if obav else "D9D9D9")
            c.font = Font(bold=True, color="000000")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.comment = Comment(("OBAVEZNO. " if obav else "Neobavezno. ") + tekst, "StockOptimizer", width=340, height=120)
            d = ws.cell(row=2, column=i)
            d.font = Font(italic=True, color="666666", size=9)
            d.alignment = Alignment(wrap_text=True, vertical="top")
            d.fill = PatternFill("solid", fgColor="F2F2F2")
            ws.column_dimensions[c.column_letter].width = 16 if i > 11 else 24
        ws.row_dimensions[1].height = 32
        ws.row_dimensions[2].height = 96
        ws.freeze_panes = "B3"
        wu = w.book["Upute"]
        wu.column_dimensions["A"].width = 30
        wu.column_dimensions["B"].width = 110
        for r in wu.iter_rows(min_row=2):
            for c in r:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in wu[1]:
            c.font = Font(bold=True)
    return buf.getvalue()


def predlozak_csv() -> bytes:
    """Prazna CSV špranca: zaglavlje i opisni redak (točka-zarez, UTF-8 s BOM-om da se u Excelu ispravno vide hrvatska slova)."""
    return ("\ufeff" + ";".join(STUPCI_PREDLOSKA) + "\n" + ";".join(f'"{c}"' for c in _opisni_redak()) + "\n").encode("utf-8")
